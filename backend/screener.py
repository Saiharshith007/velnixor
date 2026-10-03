"""Resume text extraction and per-signal scoring.

The weighted overall score is computed in the frontend so weights can change without re-parsing files.
"""
import math
import os
import re
from collections import Counter
from datetime import date

import docx
import pdfplumber

SUPPORTED = (".pdf", ".docx", ".txt", ".md")

STOPWORDS = set("""
a about above across re ll ve after again against all also am an and any are as at be because been before being below
between both but by can could did do does doing down during each etc few for from further had has have having he
her here hers him his how i if in into is it its itself just may me more most must my no nor not now of off on
once only or other our ours out over own per same shall she should so some such than that the their theirs them
then there these they this those through to too under until up upon very via was we well were what when where
which while who whom why will with within would you your yours ability able candidate candidates company day
description experience good great including job knowledge looking new plus preferred related required
requirements responsibilities role skill skills strong team teams understanding using work working year years
hiring behind build building day days own deep clear habit hands improve help join want need like love based
ideal opportunity environment fast paced passion passionate best world million millions
""".split())

_TOKEN = re.compile(r"[a-z0-9][a-z0-9+#.]*[a-z0-9+#]")


def extract_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        with pdfplumber.open(path) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    if ext == ".docx":
        d = docx.Document(path)
        # Many resume templates put content in tables.
        cells = [c.text for t in d.tables for row in t.rows for c in row.cells]
        return "\n".join([p.text for p in d.paragraphs] + cells)
    with open(path, encoding="utf-8", errors="ignore") as f:
        return f.read()


def tokens(text, surface=None):
    """Lowercased content words with crude plural folding ("systems" == "system").
    If `surface` is a dict, it records each folded token's first original spelling for display."""
    out = []
    for raw in _TOKEN.findall(text.lower()):
        if raw in STOPWORDS or raw.replace(".", "").isdigit():
            continue
        t = raw[:-1] if len(raw) > 3 and raw.endswith("s") and not raw.endswith("ss") else raw
        if surface is not None:
            surface.setdefault(t, raw)
        out.append(t)
    return out


def jd_match(jd_terms, doc_terms):
    """Share of the job description's vocabulary found in the resume, weighted by how often the JD repeats it."""
    weights = {t: 1 + math.log(c) for t, c in jd_terms.items()}
    found = sum(w for t, w in weights.items() if t in doc_terms)
    missing = sorted((t for t in weights if t not in doc_terms), key=lambda t: -weights[t])
    return round(100 * found / sum(weights.values())), missing[:12]


def term_regex(term):
    body = r"\s+".join(map(re.escape, term.split()))
    return re.compile(rf"(?<![a-z0-9]){body}(?:e?s)?(?![a-z0-9])", re.I)


def match_terms(terms, text):
    matched = [t for t in terms if term_regex(t).search(text)]
    return {"matched": matched, "missing": [t for t in terms if t not in matched]}


# --- Experience -------------------------------------------------------------------------------------------------

_STATED = [
    re.compile(r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)\.?\s+(?:of\s+)?(?:[\w-]+\s+){0,3}?(?:experience|exp\b|career)", re.I),
    re.compile(r"experience\s*[:\-–]\s*(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)", re.I),
]
_MONTHS = "jan feb mar apr may jun jul aug sep oct nov dec".split()
_POINT = r"(?:(?P<m{0}>[a-z]{{3,9}})\.?\s*,?\s*'?|(?P<n{0}>\d{{1,2}})\s*[/.-]\s*)?(?P<y{0}>(?:19|20)\d{{2}})"
_RANGE = re.compile(
    _POINT.format(1) + r"\s*(?:-|–|—|to|till|until)\s*(?:" + _POINT.format(2) + r"|(?P<now>present|current|now|date|today))",
    re.I,
)
_EDUCATION = re.compile(
    r"universit|college|school|institute|academy|bachelor|master|\bb\.?\s?tech|\bm\.?\s?tech|\bb\.?\s?sc|\bm\.?\s?sc"
    r"|\bmba\b|ph\.?d|degree|diploma|\bgpa\b|cgpa|graduat|coursework",
    re.I,
)


def _month(name, num):
    if name and name[:3].lower() in _MONTHS:
        return _MONTHS.index(name[:3].lower()) + 1
    if num and 1 <= int(num) <= 12:
        return int(num)
    return 1


def experience(text, today=None):
    """Return (years, source). Prefers an explicitly stated total, else merges employment date ranges."""
    flat = re.sub(r"\s+", " ", text)
    stated = [float(m.group(1)) for p in _STATED for m in p.finditer(flat) if float(m.group(1)) < 50]
    if stated:
        return max(stated), "stated"

    today = today or date.today()
    now = today.year * 12 + today.month - 1
    spans, lines = [], text.splitlines()
    for i, line in enumerate(lines):
        if _EDUCATION.search(line) or (i and _EDUCATION.search(lines[i - 1])):
            continue
        for m in _RANGE.finditer(line):
            start = int(m["y1"]) * 12 + _month(m["m1"], m["n1"]) - 1
            end = now if m["now"] else int(m["y2"]) * 12 + _month(m["m2"], m["n2"]) - 1
            if 1970 * 12 <= start < end <= now:
                spans.append((start, end))

    months, cur = 0, None
    for s, e in sorted(spans):  # merge overlapping jobs so concurrent roles aren't double-counted
        if cur and s <= cur[1]:
            cur[1] = max(cur[1], e)
        else:
            months += cur[1] - cur[0] if cur else 0
            cur = [s, e]
    months += cur[1] - cur[0] if cur else 0
    return (round(months / 12, 1), "dates") if months else (None, None)


# --- Contact details --------------------------------------------------------------------------------------------

_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_PHONE = re.compile(r"\+?\(?\d[\d \t().-]{7,}\d")
_LINKEDIN = re.compile(r"(?:https?://)?(?:[a-z]{2,3}\.)?linkedin\.com/in/[\w-]+/?", re.I)


def _looks_like_name(line):
    words = line.split()
    return 2 <= len(words) <= 4 and all(w[0].isupper() and all(c.isalpha() or c in ".'-" for c in w) for w in words)


def contact(text):
    email = _EMAIL.search(text)
    phone = next((p.strip() for p in _PHONE.findall(text) if 10 <= sum(c.isdigit() for c in p) <= 15), None)
    linkedin = _LINKEDIN.search(text)
    head = [line.strip() for line in text.splitlines() if line.strip()][:5]
    name = next((h for h in head if _looks_like_name(h) and not re.search(r"resume|curriculum|vitae|profile", h, re.I)), None)
    return {
        "candidate": name,
        "email": email and email.group(0),
        "phone": phone,
        "linkedin": linkedin and linkedin.group(0),
    }


# --- Pipeline ---------------------------------------------------------------------------------------------------

def screen(files, jd="", keywords=(), skills=()):
    """files: [(path, display_name)]. Yields progress events, then {"type": "done", "results": [...]}."""
    results, docs = [], []
    for i, (path, name) in enumerate(files):
        yield {"type": "progress", "done": i, "total": len(files), "file": name}
        try:
            text = extract_text(path)
            error = None if text.strip() else "No readable text. It may be a scanned image."
        except Exception as e:  # corrupt or password-protected files shouldn't sink the batch
            text, error = "", f"Could not read file ({type(e).__name__})."
        years, source = experience(text)
        results.append({
            "name": name,
            "format": os.path.splitext(name)[1][1:].upper(),
            "error": error,
            "words": len(text.split()),
            "skills": match_terms(skills, text),
            "keywords": match_terms(keywords, text),
            "experience_years": years,
            "experience_source": source,
            "text": text[:30000],
            **contact(text),
        })
        docs.append(set(tokens(text)))

    surface = {}
    jd_terms = Counter(tokens(jd, surface))
    for r, doc in zip(results, docs):
        r["jd_score"], missing = jd_match(jd_terms, doc) if jd_terms else (None, [])
        r["jd_missing"] = [surface[t] for t in missing]
    yield {"type": "done", "results": results}


if __name__ == "__main__":
    today = date(2026, 1, 1)
    assert experience("Senior engineer with 8+ years of professional experience")[0] == 8
    assert experience("Total Experience: 5 years")[0] == 5
    jobs = "Acme Corp  Jan 2018 - Dec 2020\nGlobex  Jun 2020 – Present\nXYZ University 2010 - 2014"
    assert experience(jobs, today) == (round((2025 * 12 + 12 - (2018 * 12)) / 12, 1), "dates"), experience(jobs, today)
    assert experience("Lives in a 2-bedroom flat", today) == (None, None)
    assert match_terms(["Java", "C++", "machine learning", "API"], "JavaScript, C++ and Machine\nLearning APIs") == {
        "matched": ["C++", "machine learning", "API"], "missing": ["Java"]}
    assert contact("Jane Doe\njane@doe.dev\n+1 (555) 123-4567\n2015 - 2018")["phone"] == "+1 (555) 123-4567"
    assert contact("Jane Doe\nworked 2015 - 2018")["candidate"] == "Jane Doe"
    assert contact("Résumé\nTomás Ålvarez\nSan Francisco, CA")["candidate"] == "Tomás Ålvarez"
    assert jd_match(Counter(tokens("Python Python Django AWS")), set(tokens("python django")))[0] == round(
        100 * (1 + math.log(2) + 1) / (1 + math.log(2) + 2))
    print("screener self-check ok")
